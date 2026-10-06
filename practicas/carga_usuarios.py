"""Carga masiva de usuarios desde Excel (.xlsx) o CSV, y la plantilla para llenarlos."""
import csv
import io
import unicodedata
from io import BytesIO

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Carrera, Empresa, Usuario
from .validadores import cedula_valida

COLUMNAS = ["cedula", "nombres", "apellidos", "correo", "telefono", "rol", "carrera", "empresa", "cargo"]
OBLIGATORIAS = ["cedula", "nombres", "apellidos", "correo", "rol"]

TITULOS = {
    "cedula": "CÉDULA", "nombres": "NOMBRES", "apellidos": "APELLIDOS", "correo": "CORREO",
    "telefono": "TELÉFONO", "rol": "ROL", "carrera": "CARRERA (código)",
    "empresa": "EMPRESA (RUC)", "cargo": "CARGO",
}

AZUL = "1F3B73"


def _normalizar(texto):
    """'Docente-Tutor ' -> 'DOCENTE TUTOR' (sin tildes, mayúsculas)."""
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(texto.replace("-", " ").replace("_", " ").upper().split())


# Nombres aceptados en la columna ROL
ROLES = {}
for valor, etiqueta in Usuario.Rol.choices:
    ROLES[_normalizar(valor)] = valor
    ROLES[_normalizar(etiqueta)] = valor
ROLES.update({"DOCENTE TUTOR": "TUTOR", "DOCENTE COORDINADOR": "COORDINADOR", "ADMINISTRADOR": "ADMIN",
              "TUTOR EMPRESA": "TUTOR_EMP", "TUTOR DE EMPRESA": "TUTOR_EMP"})


# ---------------------------------------------------------------------------
# Plantilla
# ---------------------------------------------------------------------------
def generar_plantilla():
    wb = Workbook()
    ws = wb.active
    ws.title = "Usuarios"
    ws.append([TITULOS[c] for c in COLUMNAS])
    ejemplo = ["0912345675", "Juan Carlos", "Pérez López", "juan.perez@istam.edu.ec", "0991234567", "ESTUDIANTE"]
    carrera_ej = Carrera.objects.first()
    ws.append(ejemplo + [carrera_ej.codigo if carrera_ej else "", "", ""])

    cab = PatternFill("solid", fgColor=AZUL)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = cab
        celda.alignment = Alignment(horizontal="center", vertical="center")
    for celda in ws[2]:
        celda.font = Font(italic=True, color="808080")
    ws.row_dimensions[1].height = 22
    for letra, ancho in zip("ABCDEFGHI", [14, 22, 22, 32, 14, 18, 18, 18, 22]):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "A2"
    # Cédula y teléfono como texto para no perder el 0 inicial
    for col in ("A", "E", "H"):
        for fila in range(2, 1001):
            ws[f"{col}{fila}"].number_format = "@"

    # Hoja con las listas válidas
    listas = wb.create_sheet("Listas")
    listas.append(["ROLES", "CÓDIGO CARRERA", "NOMBRE CARRERA", "RUC EMPRESA", "NOMBRE EMPRESA"])
    roles = ["ESTUDIANTE", "TUTOR", "TUTOR_EMP", "COORDINADOR", "RECTOR", "ADMIN"]
    carreras = list(Carrera.objects.values_list("codigo", "nombre"))
    empresas = list(Empresa.objects.values_list("ruc", "nombre"))
    for i in range(max(len(roles), len(carreras), len(empresas))):
        listas.append([
            roles[i] if i < len(roles) else None,
            carreras[i][0] if i < len(carreras) else None,
            carreras[i][1] if i < len(carreras) else None,
            empresas[i][0] if i < len(empresas) else None,
            empresas[i][1] if i < len(empresas) else None,
        ])
    for celda in listas[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = cab
    listas.column_dimensions["A"].width = 16
    listas.column_dimensions["B"].width = 18
    listas.column_dimensions["C"].width = 40
    listas.column_dimensions["D"].width = 16
    listas.column_dimensions["E"].width = 40

    # Listas desplegables
    dv_rol = DataValidation(type="list", formula1=f"=Listas!$A$2:$A${len(roles) + 1}", allow_blank=False)
    dv_rol.error, dv_rol.errorTitle = "Elija un rol de la lista.", "Rol no válido"
    ws.add_data_validation(dv_rol)
    dv_rol.add("F2:F1000")
    if carreras:
        dv_car = DataValidation(type="list", formula1=f"=Listas!$B$2:$B${len(carreras) + 1}", allow_blank=True)
        dv_car.error, dv_car.errorTitle = "Elija un código de carrera de la hoja 'Listas'.", "Carrera no válida"
        ws.add_data_validation(dv_car)
        dv_car.add("G2:G1000")

    # Hoja de instrucciones
    ins = wb.create_sheet("Instrucciones")
    textos = [
        "INSTRUCCIONES PARA LA CARGA MASIVA DE USUARIOS",
        "",
        "1. Llene una fila por usuario en la hoja 'Usuarios'. Borre la fila de ejemplo (gris).",
        "2. Columnas obligatorias: CÉDULA, NOMBRES, APELLIDOS, CORREO y ROL.",
        "3. CÉDULA: 10 dígitos, con el 0 inicial. Será el usuario para ingresar.",
        "4. La contraseña inicial de cada usuario es su propia cédula; el sistema le pedirá cambiarla al ingresar.",
        "5. ROL: ESTUDIANTE, TUTOR (docente-tutor), TUTOR_EMP (tutor empresarial), COORDINADOR, RECTOR o ADMIN.",
        "6. CARRERA: código de la carrera (ver hoja 'Listas'). Obligatoria para estudiantes.",
        "7. EMPRESA (RUC) y CARGO: solo para tutores empresariales (ver RUC en la hoja 'Listas').",
        "8. No cambie el orden ni los títulos de las columnas.",
        "9. Guarde el archivo como .xlsx y súbalo en Administración > Carga masiva.",
    ]
    for t in textos:
        ins.append([t])
    ins["A1"].font = Font(bold=True, size=13, color=AZUL)
    ins.column_dimensions["A"].width = 100

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# Lectura del archivo
# ---------------------------------------------------------------------------
def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _leer_filas(archivo):
    """Devuelve una lista de (número_de_fila, dict) a partir de .xlsx o .csv."""
    nombre = archivo.name.lower()
    if nombre.endswith(".csv"):
        contenido = archivo.read()
        for codif in ("utf-8-sig", "latin-1"):
            try:
                texto = contenido.decode(codif)
                break
            except UnicodeDecodeError:
                continue
        muestra = texto[:2000]
        separador = ";" if muestra.count(";") > muestra.count(",") else ","
        filas = list(csv.reader(io.StringIO(texto), delimiter=separador))
    else:
        wb = load_workbook(archivo, read_only=True, data_only=True)
        ws = wb["Usuarios"] if "Usuarios" in wb.sheetnames else wb.worksheets[0]
        filas = [list(f) for f in ws.iter_rows(values_only=True)]
    if not filas:
        return []

    # Identifica las columnas por su título
    encabezado = [_normalizar(c) for c in filas[0]]
    indices = {}
    for col in COLUMNAS:
        buscado = _normalizar(TITULOS[col])
        for i, titulo in enumerate(encabezado):
            if titulo == buscado or titulo.split(" (")[0] == buscado.split(" (")[0] or titulo == _normalizar(col):
                indices[col] = i
                break
    faltan = [TITULOS[c] for c in OBLIGATORIAS if c not in indices]
    if faltan:
        raise ValueError("El archivo no tiene las columnas: " + ", ".join(faltan) + ". Use la plantilla.")

    datos = []
    for n, fila in enumerate(filas[1:], start=2):
        registro = {c: _texto(fila[i]) if i < len(fila) else "" for c, i in indices.items()}
        if any(registro.values()):
            datos.append((n, registro))
    return datos


def procesar_archivo(archivo, actualizar=False):
    """Crea (o actualiza) usuarios. Devuelve un resumen con el resultado de cada fila."""
    filas = _leer_filas(archivo)
    empresas = {e.ruc: e for e in Empresa.objects.all()}
    carreras = {}
    for c in Carrera.objects.all():
        carreras[_normalizar(c.codigo)] = c
        carreras[_normalizar(c.nombre)] = c

    resultado = {"creados": 0, "actualizados": 0, "omitidos": 0, "errores": 0, "filas": []}
    vistas = set()
    for n, r in filas:
        errores = []
        cedula = r.get("cedula", "")
        if len(cedula) == 9 and cedula.isdigit():
            cedula = "0" + cedula  # Excel quitó el cero inicial
        for campo in OBLIGATORIAS:
            if not r.get(campo):
                errores.append(f"falta {TITULOS[campo].lower()}")
        if cedula and not cedula_valida(cedula):
            errores.append("cédula no válida")
        if cedula in vistas:
            errores.append("cédula repetida en el archivo")
        correo = r.get("correo", "")
        if correo:
            try:
                validate_email(correo)
            except ValidationError:
                errores.append("correo no válido")
        rol = ROLES.get(_normalizar(r.get("rol", "")))
        if r.get("rol") and not rol:
            errores.append(f"rol '{r['rol']}' no existe")
        carrera = None
        if r.get("carrera"):
            carrera = carreras.get(_normalizar(r["carrera"]))
            if not carrera:
                errores.append(f"carrera '{r['carrera']}' no existe")
        elif rol == Usuario.Rol.ESTUDIANTE:
            errores.append("el estudiante debe tener carrera")
        empresa = None
        if r.get("empresa"):
            ruc = r["empresa"]
            if len(ruc) == 12 and ruc.isdigit():
                ruc = "0" + ruc
            empresa = empresas.get(ruc)
            if not empresa:
                errores.append(f"empresa con RUC '{r['empresa']}' no existe")

        nombre = f"{r.get('nombres', '')} {r.get('apellidos', '')}".strip()
        if errores:
            resultado["errores"] += 1
            resultado["filas"].append((n, cedula, nombre, "error", "; ".join(errores)))
            continue
        vistas.add(cedula)

        existente = Usuario.objects.filter(username=cedula).first()
        if existente and not actualizar:
            resultado["omitidos"] += 1
            resultado["filas"].append((n, cedula, nombre, "omitido", "ya existe"))
            continue

        with transaction.atomic():
            u = existente or Usuario(cedula=cedula)
            u.first_name = r["nombres"].title()
            u.last_name = r["apellidos"].title()
            u.email = correo.lower()
            u.telefono = r.get("telefono", "")
            u.rol = rol
            u.carrera = carrera
            u.empresa = empresa
            u.cargo = r.get("cargo", "")
            u.is_staff = rol == Usuario.Rol.ADMIN
            if not existente:
                u.set_password(cedula)
                u.debe_cambiar_clave = True
            u.save()
        if existente:
            resultado["actualizados"] += 1
            resultado["filas"].append((n, cedula, nombre, "actualizado", ""))
        else:
            resultado["creados"] += 1
            resultado["filas"].append((n, cedula, nombre, "creado", ""))
    return resultado
